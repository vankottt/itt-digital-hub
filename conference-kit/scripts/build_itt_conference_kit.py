#!/usr/bin/env python3
"""Build the ITT Digital Hub conference kit from selected diva-e MARKETING files."""

from __future__ import annotations

import math
import shutil
import urllib.request
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
LIB = Path(
    "/Users/ivan.todorov/Documents/Projects/diva-e/diva-e MARKETING"
    "/Marketing Database Offline/01_Library"
)
KIT = ROOT / "conference-kit"
SRC_OUT = KIT / "01-selected-sources"
BRANDED = KIT / "02-itt-branded"
BRAND = BRANDED / "brand"
FONTS = BRANDED / "fonts"
PPT_OUT = BRANDED / "powerpoint"
DIGITAL = BRANDED / "digital"
PRINT = BRANDED / "print"
ICONS_OUT = BRANDED / "icons"
PREVIEWS = BRANDED / "previews"

# Website tokens — src/app/globals.css
PAPER = (237, 240, 251)
PAPER_2 = (225, 231, 246)
INK = (4, 14, 49)
INK_2 = (58, 70, 96)
INK_3 = (91, 103, 128)
LINE = (220, 226, 238)
LINE_STRONG = (180, 189, 201)
MARINE = (4, 14, 49)
MARINE_2 = (10, 24, 80)
MARINE_TINT = (232, 237, 255)
SPRUCE = (30, 107, 88)
SIGNAL = (0, 44, 255)
SIGNAL_2 = (27, 61, 255)
ON_DARK = (247, 248, 253)
ON_DARK_MUTED = (154, 168, 196)
WHITE = (255, 255, 255)

COLOR_MAP = {
    "081235": "040E31",
    "1B31FF": "002CFF",
    "11179A": "0018B0",
    "8490FD": "1B3DFF",
    "A7FF71": "1E6B58",
    "9C73F3": "1A5B4B",
}

FONT_REPLACEMENTS = [
    ("GT Planar Italic 30", "IBM Plex Sans"),
    ("GT Planar Italic", "IBM Plex Sans"),
    ("GT Planar", "IBM Plex Sans"),
    ("Inter", "IBM Plex Sans"),
]

COPIES = [
    (
        "PPT Master & Examples/ppt_Master_Logo Endorsement.potx",
        "powerpoint/ppt_Master_Logo Endorsement.potx",
    ),
    (
        "PPT Master & Examples/Icon Collection.potx",
        "powerpoint/Icon Collection.potx",
    ),
    (
        "PPT Master & Examples/Power_Point_Training.pptx",
        "powerpoint/Power_Point_Training.pptx",
    ),
    (
        "PPT Master & Examples/How-to-change the logo in your ppt.webm",
        "powerpoint/How-to-change the logo in your ppt.webm",
    ),
    (
        "TEMPLATES/Poster Templates/Poster_1.potx",
        "posters/Poster_1.potx",
    ),
    (
        "TEMPLATES/Poster Templates/Poster_2.potx",
        "posters/Poster_2.potx",
    ),
    (
        "TEMPLATES/Poster Templates/Poster_3.potx",
        "posters/Poster_3.potx",
    ),
    (
        "TEMPLATES/Poster Templates/Poster_1_Druck.potx",
        "posters/Poster_1_Druck.potx",
    ),
    (
        "TEMPLATES/Poster Templates/Poster_2_Druck.potx",
        "posters/Poster_2_Druck.potx",
    ),
    (
        "TEMPLATES/Poster Templates/Poster_3_Druck.potx",
        "posters/Poster_3_Druck.potx",
    ),
    ("LinkedIn Header/LinkedIn Header_ New Logo.png", "linkedin/LinkedIn Header_ New Logo.png"),
    ("Letter templates/Word_Template.dotx", "letter/Word_Template.dotx"),
    (
        "Letter templates/International/Word_template_Bulgaria_Sofia.dotx",
        "letter/Word_template_Bulgaria_Sofia.dotx",
    ),
    (
        "MS Teams Backgrounds/1920x1080/1920 x 1080 px.jpg",
        "teams/1920 x 1080 px.jpg",
    ),
    ("BRAND/Event Equipment/Background_Wall.pdf", "event/Background_Wall.pdf"),
    ("BRAND/Event Equipment/Badge_A_03_cut.pdf", "event/Badge_A_03_cut.pdf"),
    ("BRAND/Event Equipment/Badge_B_03_cut.pdf", "event/Badge_B_03_cut.pdf"),
    (
        "BRAND/Event Equipment/diva-e_VorlagePoster_druck_02.pdf",
        "event/diva-e_VorlagePoster_druck_02.pdf",
    ),
    (
        "BRAND/Event Equipment/RollUp Trio Recruiting Events.png",
        "event/RollUp Trio Recruiting Events.png",
    ),
    (
        "BRAND/Event Equipment/diva-e_moderationkarten_25_01.pdf",
        "event/diva-e_moderationkarten_25_01.pdf",
    ),
    (
        "BRAND/Event Equipment/Lanyard (Print File) 02 (1).pdf",
        "event/Lanyard (Print File) 02 (1).pdf",
    ),
    (
        "PPT Master & Examples/diva-e Folienbibliothek Sales/diva-e AI/Mission.pptx",
        "structure-examples/Mission.pptx",
    ),
    (
        "PPT Master & Examples/diva-e Folienbibliothek Sales/diva-e AI/Why_AI_Solutions.pptx",
        "structure-examples/Why_AI_Solutions.pptx",
    ),
    (
        "PPT Master & Examples/diva-e Folienbibliothek Sales/diva-e AI/Consulting_Approach.pptx",
        "structure-examples/Consulting_Approach.pptx",
    ),
    (
        "PPT Master & Examples/diva-e Folienbibliothek Sales/diva-e AI/Field_of_action.pptx",
        "structure-examples/Field_of_action.pptx",
    ),
    (
        "PPT Master & Examples/diva-e Folienbibliothek Sales/diva-e AI/Next_Steps.pptx",
        "structure-examples/Next_Steps.pptx",
    ),
]


def rgb(c: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*c)


def ensure_dirs() -> None:
    for p in (SRC_OUT, BRAND, FONTS, PPT_OUT, DIGITAL, PRINT, ICONS_OUT, PREVIEWS, BRANDED / "letter"):
        p.mkdir(parents=True, exist_ok=True)


def copy_selected() -> None:
    for rel_src, rel_dst in COPIES:
        src = LIB / rel_src
        dst = SRC_OUT / rel_dst
        if not src.exists():
            print(f"MISSING {src}")
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
        print(f"copied {rel_dst}")

    icon_src = LIB / "PPT Master & Examples/icons"
    icon_dst = SRC_OUT / "icons"
    if icon_src.exists():
        if icon_dst.exists():
            shutil.rmtree(icon_dst)
        shutil.copytree(icon_src, icon_dst)
        print("copied icons/")


def download_fonts() -> dict[str, Path]:
    files = {
        "regular": (
            "IBMPlexSans-Regular.ttf",
            "https://cdn.jsdelivr.net/fontsource/fonts/ibm-plex-sans@5.2.5/latin-400-normal.ttf",
        ),
        "medium": (
            "IBMPlexSans-Medium.ttf",
            "https://cdn.jsdelivr.net/fontsource/fonts/ibm-plex-sans@5.2.5/latin-500-normal.ttf",
        ),
        "semibold": (
            "IBMPlexSans-SemiBold.ttf",
            "https://cdn.jsdelivr.net/fontsource/fonts/ibm-plex-sans@5.2.5/latin-600-normal.ttf",
        ),
    }
    out: dict[str, Path] = {}
    for key, (name, url) in files.items():
        path = FONTS / name
        if not path.exists() or path.stat().st_size < 1000:
            try:
                urllib.request.urlretrieve(url, path)
            except Exception as exc:  # noqa: BLE001
                print(f"font download failed {name}: {exc}")
                continue
        out[key] = path
    return out


def font(fonts: dict[str, Path], key: str, size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = fonts.get(key) or fonts.get("regular")
    if path and path.exists():
        return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def radial_glow(size: tuple[int, int], color: tuple[int, int, int], strength: float = 0.55) -> Image.Image:
    w, h = size
    src_w, src_h = 280, 280
    glow = Image.new("RGBA", (src_w, src_h), (0, 0, 0, 0))
    px = glow.load()
    cx, cy = (src_w - 1) / 2, (src_h - 1) / 2
    max_d = math.hypot(cx, cy)
    for y in range(src_h):
        for x in range(src_w):
            d = math.hypot(x - cx, y - cy) / max_d
            a = int(max(0.0, 1.0 - d) ** 2.1 * 255 * strength)
            px[x, y] = (*color, a)
    return glow.resize((w, h), Image.Resampling.LANCZOS)


def hero_atmosphere(w: int = 1920, h: int = 1080) -> Image.Image:
    img = Image.new("RGB", (w, h), MARINE)
    overlay = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    g1 = radial_glow((int(w * 0.95), int(h * 0.85)), SIGNAL, 0.62)
    overlay.alpha_composite(g1, (int(w * 0.28), int(h * -0.08)))
    g2 = radial_glow((int(w * 0.55), int(h * 0.55)), MARINE_2, 0.9)
    overlay.alpha_composite(g2, (int(w * -0.08), int(h * 0.52)))
    return Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")


def paper_surface(w: int = 1920, h: int = 1080, grid: bool = True) -> Image.Image:
    img = Image.new("RGB", (w, h), PAPER)
    draw = ImageDraw.Draw(img)
    if grid:
        step = 24
        for y in range(0, h, step):
            for x in range(0, w, step):
                draw.point((x, y), fill=LINE_STRONG)
    return img


def fit_logo(src: Path, size: tuple[int, int], pad: float = 0.10) -> Image.Image:
    canvas = Image.new("RGBA", size, (0, 0, 0, 0))
    logo = Image.open(src).convert("RGBA")
    box = logo.getbbox()
    if box:
        logo = logo.crop(box)
    max_w = max(1, int(size[0] * (1 - 2 * pad)))
    max_h = max(1, int(size[1] * (1 - 2 * pad)))
    logo.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)
    x = (size[0] - logo.width) // 2
    y = (size[1] - logo.height) // 2
    canvas.paste(logo, (x, y), logo)
    return canvas


def average_opaque(im: Image.Image) -> tuple[float, float, float] | None:
    rgba = im.convert("RGBA")
    pixels = [p for p in rgba.getdata() if p[3] > 24]
    if not pixels:
        return None
    n = len(pixels)
    return (sum(p[0] for p in pixels) / n, sum(p[1] for p in pixels) / n, sum(p[2] for p in pixels) / n)


def recolor_svg_text(text: str) -> str:
    for src, dst in COLOR_MAP.items():
        text = text.replace(src, dst).replace(src.lower(), dst.lower())
    return text


def rewrite_text_files(root: Path) -> None:
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".xml", ".rels", ".svg"}:
            continue
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            continue
        original = text
        for src, dst in COLOR_MAP.items():
            text = text.replace(src, dst).replace(src.lower(), dst.lower())
        for old, new in FONT_REPLACEMENTS:
            text = text.replace(old, new)
        text = text.replace('clrScheme name="diva-e 2023"', 'clrScheme name="ITT Digital Hub"')
        text = text.replace('fontScheme name="diva-e"', 'fontScheme name="ITT"')
        text = text.replace("info@diva-e.com", "ittdigitalhub.uk")
        text = text.replace("www.diva-e.com", "ittdigitalhub.uk")
        text = text.replace("diva-e.com", "ittdigitalhub.uk")
        text = text.replace("diva-e Digital Value Excellence GmbH", "ITT Digital Hub")
        text = text.replace("diva-e CONCLUSION", "ITT Digital Hub")
        text = text.replace("diva-e Conclusion", "ITT Digital Hub")
        text = text.replace("St.-Martin-Straße 72, 81541 München", "ittdigitalhub.uk")
        text = text.replace("St.-Martin-Strasse 72, 81541 Muenchen", "ittdigitalhub.uk")
        text = text.replace("diva-e", "ITT Digital Hub")
        # Repair strings already partly rewritten on a previous pass
        text = text.replace("info@ITT Digital Hub.com", "ittdigitalhub.uk")
        text = text.replace("ITT Digital Hub.com", "ittdigitalhub.uk")
        text = text.replace(
            "ITT Digital Hub Digital Value\nExcellence GmbH",
            "ITT Digital Hub",
        )
        text = text.replace("ITT Digital Hub Digital Value Excellence GmbH", "ITT Digital Hub")
        text = text.replace("ITT Digital Hub Digital Value", "ITT Digital Hub")
        text = text.replace("Digital Value Excellence GmbH", "Applied AI Consultancy")
        text = text.replace("<a:t> Digital Value</a:t>", "<a:t></a:t>")
        text = text.replace("<a:t>Excellence GmbH</a:t>", "<a:t>Applied AI Consultancy</a:t>")
        text = text.replace("<a:t>@ittdigitalhub.uk</a:t>", "<a:t></a:t>")
        text = text.replace("<a:t>@diva-e.com</a:t>", "<a:t></a:t>")
        text = text.replace("Excellence GmbH", "Applied AI Consultancy")
        text = text.replace("IBM Plex Sans Retalic 30", "IBM Plex Sans")
        text = text.replace("IBM Plex Sans Italic 30", "IBM Plex Sans")
        text = text.replace("Kapitelthema", "Section")
        if text != original:
            path.write_text(text, encoding="utf-8")


def replace_media(unpacked: Path, lockup_light: Path, lockup_dark: Path) -> None:
    media_dirs = [p for p in unpacked.rglob("media") if p.is_dir()]
    glow = radial_glow((1046, 1048), SIGNAL, 0.7)
    dark_panel = hero_atmosphere(2222, 2500)

    for media in media_dirs:
        for path in list(media.iterdir()):
            suffix = path.suffix.lower()
            if suffix == ".svg":
                text = path.read_text(encoding="utf-8", errors="ignore")
                if path.stat().st_size > 2000 or "diva" in text.lower():
                    avg = None
                    if "#fff" in text.lower() or "ffffff" in text.lower() or "#f7f8fd" in text.lower():
                        src = lockup_dark
                    else:
                        src = lockup_light
                    png = fit_logo(src, (1192, 736), pad=0.06)
                    png_path = path.with_suffix(".png")
                    png.save(png_path)
                    for rel in unpacked.rglob("*"):
                        if rel.suffix.lower() in {".xml", ".rels"}:
                            t = rel.read_text(encoding="utf-8", errors="ignore")
                            n = t.replace(path.name, png_path.name)
                            if n != t:
                                rel.write_text(n, encoding="utf-8")
                    path.unlink()
                else:
                    path.write_text(recolor_svg_text(text), encoding="utf-8")
                continue
            if suffix != ".png":
                continue
            im = Image.open(path).convert("RGBA")
            avg = average_opaque(im)
            if avg is None:
                continue
            r, g, b = avg
            # Soft blue orb / glow
            if abs(r - g) < 40 and b > 200 and im.size[0] <= 1200:
                glow.resize(im.size, Image.Resampling.LANCZOS).save(path)
                continue
            # Dark atmospheric panels
            if r < 40 and g < 40 and b < 120 and im.size[0] >= 1800:
                dark_panel.resize(im.size, Image.Resampling.LANCZOS).save(path)
                continue
            # White / paper lockup (used on dark slides)
            if r > 200 and g > 200 and b > 220:
                fit_logo(lockup_dark, im.size, pad=0.08).save(path)
                continue
            # Navy lockup (used on light slides)
            if r < 40 and g < 40 and b < 80:
                fit_logo(lockup_light, im.size, pad=0.08).save(path)
                continue
            # Blue decorative waves — shift toward ITT signal
            if b > 180 and r < 80:
                data = list(im.getdata())
                shifted = []
                for px in data:
                    if px[3] < 12:
                        shifted.append(px)
                        continue
                    nr = int(px[0] * 0.15)
                    ng = int(px[1] * 0.85)
                    nb = 255 if px[2] > 80 else px[2]
                    shifted.append((nr, min(61, ng), min(255, max(44, nb)), px[3]))
                im.putdata(shifted)
                im.save(path)


def pack_office(src_dir: Path, dest: Path, as_pptx: bool = False) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    ct = src_dir / "[Content_Types].xml"
    if as_pptx and ct.exists():
        text = ct.read_text(encoding="utf-8")
        text = text.replace(
            "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
        )
        ct.write_text(text, encoding="utf-8")
    with zipfile.ZipFile(dest, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        files = [p for p in src_dir.rglob("*") if p.is_file()]
        files.sort(key=lambda p: (p.name != "[Content_Types].xml", str(p)))
        for path in files:
            zf.write(path, path.relative_to(src_dir).as_posix())


def retheme_office(src: Path, dest: Path, lockup_light: Path, lockup_dark: Path, as_pptx: bool = False) -> None:
    tmp = dest.parent / f".tmp_{dest.stem}"
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir(parents=True)
    with zipfile.ZipFile(src) as zf:
        zf.extractall(tmp)
    rewrite_text_files(tmp)
    replace_media(tmp, lockup_light, lockup_dark)
    pack_office(tmp, dest, as_pptx=as_pptx)
    shutil.rmtree(tmp)
    print(f"rethemed {dest.name}")


def recolor_icons() -> None:
    src = SRC_OUT / "icons"
    if not src.exists():
        return
    if ICONS_OUT.exists():
        shutil.rmtree(ICONS_OUT)
    shutil.copytree(src, ICONS_OUT)
    for path in ICONS_OUT.rglob("*.svg"):
        path.write_text(recolor_svg_text(path.read_text(encoding="utf-8", errors="ignore")), encoding="utf-8")
    print("recolored icons")


def draw_label(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, fonts: dict[str, Path], dark: bool = False) -> None:
    draw.text(xy, text.upper(), font=font(fonts, "medium", 22), fill=ON_DARK_MUTED if dark else INK_3)


def generate_digital_assets(fonts: dict[str, Path]) -> dict[str, Path]:
    lockup_light = BRAND / "itt-lockup-compact.png"
    lockup_dark = BRAND / "itt-lockup-compact-on-dark.png"
    mark_dark = BRAND / "itt-mark-on-dark.png"

    hero = hero_atmosphere()
    paper = paper_surface()
    hero.save(DIGITAL / "bg-hero-1920x1080.png")
    paper.save(DIGITAL / "bg-paper-1920x1080.png")

    # LinkedIn 1584 × 396
    li = hero_atmosphere(1584, 396)
    li_draw = ImageDraw.Draw(li)
    logo = Image.open(lockup_dark).convert("RGBA")
    logo.thumbnail((520, 110), Image.Resampling.LANCZOS)
    li.paste(logo, (72, 78), logo)
    li_draw.text((72, 210), "APPLIED AI CONSULTANCY", font=font(fonts, "medium", 22), fill=ON_DARK_MUTED)
    li_draw.text((72, 248), "From complex workflows to working AI systems.", font=font(fonts, "semibold", 28), fill=ON_DARK)
    li_draw.text((72, 330), "ittdigitalhub.uk", font=font(fonts, "regular", 20), fill=ON_DARK_MUTED)
    bar = Image.new("RGB", (72, 8), SIGNAL)
    li.paste(bar, (72, 188))
    li.save(DIGITAL / "linkedin-header-1584x396.png")

    # Teams 1920 × 1080
    teams = hero.copy()
    tdraw = ImageDraw.Draw(teams)
    logo2 = Image.open(lockup_dark).convert("RGBA")
    logo2.thumbnail((420, 90), Image.Resampling.LANCZOS)
    teams.paste(logo2, (80, 80), logo2)
    tdraw.text((80, 980), "ittdigitalhub.uk", font=font(fonts, "regular", 28), fill=ON_DARK_MUTED)
    teams.save(DIGITAL / "teams-background-1920x1080.png")

    # Title card preview
    title = hero.copy()
    td = ImageDraw.Draw(title)
    logo3 = Image.open(lockup_dark).convert("RGBA")
    logo3.thumbnail((360, 76), Image.Resampling.LANCZOS)
    title.paste(logo3, (96, 88), logo3)
    td.text((96, 320), "APPLIED AI CONSULTANCY", font=font(fonts, "medium", 24), fill=ON_DARK_MUTED)
    td.text((96, 372), "From complex workflows", font=font(fonts, "regular", 72), fill=ON_DARK)
    td.text((96, 458), "to working AI systems.", font=font(fonts, "regular", 72), fill=ON_DARK)
    td.rectangle((96, 560, 168, 568), fill=SIGNAL)
    td.text((96, 600), "ITT Digital Hub  ·  ittdigitalhub.uk", font=font(fonts, "regular", 28), fill=ON_DARK_MUTED)
    title.save(PREVIEWS / "01-title.png")

    # Roll-up preview 1080 × 1920
    roll = hero_atmosphere(1080, 1920)
    rd = ImageDraw.Draw(roll)
    rlogo = Image.open(lockup_dark).convert("RGBA")
    rlogo.thumbnail((720, 150), Image.Resampling.LANCZOS)
    roll.paste(rlogo, ((1080 - rlogo.width) // 2, 220), rlogo)
    rd.text((90, 520), "APPLIED AI CONSULTANCY", font=font(fonts, "medium", 28), fill=ON_DARK_MUTED)
    rd.text((90, 620), "From complex", font=font(fonts, "regular", 64), fill=ON_DARK)
    rd.text((90, 700), "workflows to", font=font(fonts, "regular", 64), fill=ON_DARK)
    rd.text((90, 780), "working AI", font=font(fonts, "regular", 64), fill=ON_DARK)
    rd.text((90, 860), "systems.", font=font(fonts, "regular", 64), fill=ON_DARK)
    rd.rectangle((90, 980, 180, 988), fill=SIGNAL)
    rd.text((90, 1040), "Process first. AI where it\nmakes sense. Build what\nwe recommend.", font=font(fonts, "regular", 32), fill=ON_DARK_MUTED)
    rd.text((90, 1760), "ittdigitalhub.uk", font=font(fonts, "medium", 32), fill=ON_DARK)
    roll.save(PRINT / "rollup-preview-1080x1920.png")

    # Badge 85×54 mm at 300 dpi ≈ 1004 × 638
    badge = Image.new("RGB", (1004, 638), MARINE)
    glow = radial_glow((700, 500), SIGNAL, 0.5)
    badge = Image.alpha_composite(badge.convert("RGBA"), Image.new("RGBA", badge.size, (0, 0, 0, 0)))
    badge.alpha_composite(glow, (400, -40))
    badge = badge.convert("RGB")
    bd = ImageDraw.Draw(badge)
    mlogo = Image.open(mark_dark).convert("RGBA")
    mlogo.thumbnail((180, 110), Image.Resampling.LANCZOS)
    badge.paste(mlogo, (56, 48), mlogo)
    bd.text((56, 200), "ITT Digital Hub", font=font(fonts, "semibold", 42), fill=ON_DARK)
    bd.text((56, 260), "Applied AI Consultancy", font=font(fonts, "regular", 28), fill=ON_DARK_MUTED)
    bd.rectangle((56, 330, 160, 338), fill=SIGNAL)
    bd.text((56, 380), "Name", font=font(fonts, "regular", 36), fill=ON_DARK)
    bd.text((56, 430), "Role / company", font=font(fonts, "regular", 24), fill=ON_DARK_MUTED)
    bd.text((56, 560), "ittdigitalhub.uk", font=font(fonts, "medium", 22), fill=ON_DARK_MUTED)
    badge.save(PRINT / "badge-85x54mm.png")

    # One-pager A4 1240 × 1754
    page = paper_surface(1240, 1754, grid=False)
    pd = ImageDraw.Draw(page)
    plogo = Image.open(lockup_light).convert("RGBA")
    plogo.thumbnail((420, 90), Image.Resampling.LANCZOS)
    page.paste(plogo, (80, 70), plogo)
    pd.text((80, 200), "APPLIED AI CONSULTANCY", font=font(fonts, "medium", 20), fill=INK_3)
    pd.text((80, 250), "From complex workflows\nto working AI systems.", font=font(fonts, "regular", 44), fill=INK)
    pd.rectangle((80, 380, 150, 388), fill=SIGNAL)
    body = (
        "ITT Digital Hub combines business process expertise and\n"
        "hands-on software engineering to design and build AI\n"
        "solutions around real operations, existing systems and data.\n\n"
        "Process first. AI where it makes sense. Build what we recommend."
    )
    pd.text((80, 430), body, font=font(fonts, "regular", 26), fill=INK_2)
    cards = [
        ("01", "Understand", "Workflow, people, systems, data and the actual business problem — before choosing a technology."),
        ("02", "Design", "The right combination of AI, automation, software, integration, data and human control."),
        ("03", "Build", "Hands-on implementation. The same senior people stay involved in building it."),
    ]
    y = 720
    for code, title, body in cards:
        pd.rounded_rectangle((80, y, 1160, y + 210), radius=28, fill=WHITE)
        pd.text((110, y + 28), code, font=font(fonts, "medium", 18), fill=INK_3)
        pd.text((110, y + 62), title, font=font(fonts, "semibold", 30), fill=INK)
        pd.text((110, y + 112), body, font=font(fonts, "regular", 22), fill=INK_2)
        y += 230
    pd.text((80, 1650), "ittdigitalhub.uk", font=font(fonts, "medium", 22), fill=SIGNAL)
    page.save(PRINT / "one-pager-a4.png")

    return {
        "hero": DIGITAL / "bg-hero-1920x1080.png",
        "paper": DIGITAL / "bg-paper-1920x1080.png",
        "linkedin": DIGITAL / "linkedin-header-1584x396.png",
    }


def set_run(run, text: str, size: int, color: tuple[int, int, int], bold: bool = False) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = "IBM Plex Sans"
    run.font.italic = False


def add_bg(slide, path: Path) -> None:
    slide.shapes.add_picture(str(path), Emu(0), Emu(0), width=Inches(13.333), height=Inches(7.5))


def add_lockup(slide, path: Path, left: float, top: float, width: float) -> None:
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def add_text(slide, l, t, w, h, lines: list[tuple[str, int, tuple[int, int, int], bool]], align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        run = p.add_run()
        set_run(run, text, size, color, bold)
    return box


def add_card(slide, l, t, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(WHITE)
    shape.line.fill.background()
    # tighter radius via adj if present
    try:
        shape.adjustments[0] = 0.12
    except Exception:
        pass
    return shape


def add_signal_bar(slide, l, t, w=0.55, h=0.06):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(SIGNAL)
    shape.line.fill.background()
    return shape


def build_starter_deck(bg_hero: Path, bg_paper: Path) -> Path:
    lockup_light = BRAND / "itt-lockup-compact.png"
    lockup_dark = BRAND / "itt-lockup-compact-on-dark.png"
    portrait = ROOT / "public/images/team/ivan-todorov-portrait-v2.jpg"
    portrait_tomchev = ROOT / "public/images/team/ivan-tomchev-portrait-v2.jpg"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_hero)
    add_lockup(s, lockup_dark, 0.7, 0.45, 2.6)
    add_text(s, 0.7, 2.15, 11, 0.4, [("APPLIED AI CONSULTANCY", 13, ON_DARK_MUTED, True)])
    add_text(
        s,
        0.7,
        2.55,
        12,
        2.2,
        [
            ("From complex workflows", 40, ON_DARK, False),
            ("to working AI systems.", 40, ON_DARK, False),
        ],
    )
    add_signal_bar(s, 0.7, 4.85)
    add_text(s, 0.7, 5.15, 10, 0.8, [("ITT Digital Hub  ·  ittdigitalhub.uk", 16, ON_DARK_MUTED, False)])
    add_text(s, 0.7, 6.75, 10, 0.3, [("Conference introduction", 12, ON_DARK_MUTED, False)])

    # 2 Positioning
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_paper)
    add_lockup(s, lockup_light, 0.7, 0.35, 2.3)
    add_text(s, 0.7, 1.3, 11, 0.35, [("POSITIONING", 12, INK_3, True)])
    add_text(s, 0.7, 1.7, 11.5, 1.6, [("Business understanding × systems thinking × AI engineering.", 28, INK, False)])
    add_signal_bar(s, 0.7, 3.4)
    add_text(
        s,
        0.7,
        3.7,
        11.5,
        2.4,
        [
            (
                "We combine business process expertise and hands-on software engineering to design and build AI solutions around real operations, existing systems and data.",
                18,
                INK_2,
                False,
            ),
            ("AI is a capability, not the product.", 18, INK, True),
        ],
    )

    # 3 Judgement
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_hero)
    add_text(s, 0.7, 1.5, 11, 0.35, [("JUDGEMENT", 12, ON_DARK_MUTED, True)])
    add_text(s, 0.7, 2.0, 12, 1.4, [("AI isn’t always the answer.", 36, ON_DARK, False)])
    add_signal_bar(s, 0.7, 3.55)
    add_text(
        s,
        0.7,
        3.9,
        11.5,
        2.2,
        [
            (
                "We start with the process, not the model. Depending on the problem, the right architecture may involve AI agents, deterministic automation, integrations, conventional software, existing enterprise tools, local AI — or a combination.",
                18,
                ON_DARK_MUTED,
                False,
            )
        ],
    )

    # 4 Problems
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_paper)
    add_text(s, 0.7, 0.4, 11, 0.3, [("WHAT WE SOLVE", 12, INK_3, True)])
    add_text(s, 0.7, 0.75, 12, 0.7, [("Problems, not a service catalogue.", 28, INK, False)])
    problems = [
        ("01", "Fragmented workflows", "Work split across people, systems, email, spreadsheets, manual handoffs and disconnected applications."),
        ("02", "Knowledge-heavy work", "Skilled people spend too long finding, interpreting, preparing, validating and summarising information."),
        ("03", "Complex operations", "Settings where software, AI, data, operational context, physical systems and human decisions have to work together."),
    ]
    for i, (code, title, body) in enumerate(problems):
        left = 0.7 + i * 4.1
        add_card(s, left, 2.0, 3.85, 4.5)
        add_text(s, left + 0.3, 2.25, 3.3, 0.3, [(code, 12, INK_3, True)])
        add_text(s, left + 0.3, 2.7, 3.3, 1.1, [(title, 20, INK, False)])
        add_text(s, left + 0.3, 3.9, 3.3, 2.2, [(body, 14, INK_2, False)])

    # 5 Approach
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_paper)
    add_text(s, 0.7, 0.4, 11, 0.3, [("HOW WE WORK", 12, INK_3, True)])
    add_text(s, 0.7, 0.75, 12, 0.7, [("Understand. Design. Build.", 28, INK, False)])
    steps = [
        ("01", "Understand", "Workflow, people, systems, data, constraints and the actual business problem — before choosing a technology."),
        ("02", "Design", "The right combination of AI, automation, software, integration, data and human control."),
        ("03", "Build", "Hands-on implementation. The same senior people who understood the problem stay involved in building it."),
    ]
    for i, (code, title, body) in enumerate(steps):
        left = 0.7 + i * 4.1
        add_card(s, left, 2.0, 3.85, 4.5)
        add_text(s, left + 0.3, 2.25, 3.3, 0.3, [(code, 12, SIGNAL, True)])
        add_text(s, left + 0.3, 2.7, 3.3, 0.8, [(title, 22, INK, False)])
        add_text(s, left + 0.3, 3.7, 3.3, 2.3, [(body, 14, INK_2, False)])

    # 6 Work
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_paper)
    add_text(s, 0.7, 0.4, 11, 0.3, [("WORK", 12, INK_3, True)])
    add_text(s, 0.7, 0.75, 12, 0.6, [("Our stories", 28, INK, False)])
    stories = [
        ("AI-Assisted Solar Operations", "Applied AI around solar infrastructure: monitoring and operational support where they are verified. Status: in development."),
        ("Local AI Orchestration", "Coordinating specialised local models, tools and services. Internal R&D / prototype — not a product being sold."),
        ("ATN Warranty Portal", "Warranty registration and operator workflow. Working software, without forcing AI into the story. warranty.atneu.com"),
    ]
    for i, (title, body) in enumerate(stories):
        top = 1.7 + i * 1.75
        add_card(s, 0.7, top, 11.9, 1.55)
        add_text(s, 1.0, top + 0.2, 11.3, 0.4, [(title, 18, INK, False)])
        add_text(s, 1.0, top + 0.7, 11.3, 0.65, [(body, 14, INK_2, False)])

    # 7 Team
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_paper)
    add_text(s, 0.7, 0.4, 11, 0.3, [("TEAM", 12, INK_3, True)])
    add_text(s, 0.7, 0.75, 12, 0.6, [("Two complementary specialists.", 28, INK, False)])
    add_card(s, 0.7, 1.8, 5.8, 4.8)
    add_card(s, 6.8, 1.8, 5.8, 4.8)
    if portrait.exists():
        s.shapes.add_picture(str(portrait), Inches(1.0), Inches(2.05), width=Inches(1.5), height=Inches(1.5))
    add_text(s, 2.7, 2.15, 3.5, 0.4, [("Ivan Todorov", 20, INK, False)])
    add_text(s, 2.7, 2.55, 3.5, 0.7, [("AI Strategy & Business Transformation Consultant", 13, INK_2, False)])
    add_text(
        s,
        1.0,
        3.8,
        5.2,
        2.3,
        [
            ("I work at the intersection of business strategy, processes and applied AI.", 14, INK_2, False)
        ],
    )
    if portrait_tomchev.exists():
        s.shapes.add_picture(str(portrait_tomchev), Inches(7.1), Inches(2.05), width=Inches(1.5), height=Inches(1.5))
    add_text(s, 8.8, 2.15, 3.5, 0.4, [("Ivan Tomchev", 20, INK, False)])
    add_text(s, 8.8, 2.55, 3.5, 0.7, [("AI Systems Architect & Software Engineer", 13, INK_2, False)])
    add_text(
        s,
        7.1,
        3.8,
        5.2,
        2.3,
        [
            ("I design and build software and AI systems for complex operational environments.", 14, INK_2, False)
        ],
    )

    # 8 Partners
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_paper)
    add_text(s, 0.7, 0.4, 11, 0.3, [("PARTNERS", 12, INK_3, True)])
    add_text(s, 0.7, 0.75, 12, 0.8, [("Organisations with real operations.", 28, INK, False)])
    add_text(
        s,
        0.7,
        1.7,
        12,
        0.8,
        [("Retail, manufacturing, logistics and public institutions — as shown on ittdigitalhub.uk. Not claimed as clients unless confirmed.", 16, INK_2, False)],
    )
    names = ["ATN", "Merkanto", "Tomchevi", "UACG", "UNWE", "MG Klima", "Nakra"]
    for i, name in enumerate(names):
        left = 0.7 + (i % 4) * 3.1
        top = 2.8 + (i // 4) * 1.8
        add_card(s, left, top, 2.9, 1.5)
        add_text(s, left + 0.2, top + 0.5, 2.5, 0.5, [(name, 16, INK, False)], align=PP_ALIGN.CENTER)

    # 9 Contact
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_hero)
    add_lockup(s, lockup_dark, 0.7, 0.45, 2.6)
    add_text(s, 0.7, 2.2, 11, 0.35, [("CONTACT", 12, ON_DARK_MUTED, True)])
    add_text(s, 0.7, 2.65, 12, 1.4, [("Have a problem worth solving?", 32, ON_DARK, False)])
    add_signal_bar(s, 0.7, 4.2)
    add_text(
        s,
        0.7,
        4.55,
        11.5,
        1.6,
        [
            ("If you have a process, operation or system where AI might make sense — show us the problem.", 18, ON_DARK_MUTED, False),
            ("ittdigitalhub.uk", 20, ON_DARK, True),
        ],
    )

    # 10 Section divider template
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_hero)
    add_text(s, 0.7, 3.0, 11, 0.35, [("SECTION LABEL", 12, ON_DARK_MUTED, True)])
    add_text(s, 0.7, 3.4, 12, 1.2, [("Section title", 40, ON_DARK, False)])

    # 11 Quote template
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_paper)
    add_signal_bar(s, 0.7, 2.4, 0.7, 0.07)
    add_text(s, 0.7, 2.7, 12, 2.2, [("A short statement that carries the slide.", 32, INK, False)])
    add_text(s, 0.7, 5.2, 10, 0.4, [("Optional attribution", 14, INK_3, False)])

    # 12 Two column template
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_paper)
    add_text(s, 0.7, 0.45, 11, 0.3, [("LABEL", 12, INK_3, True)])
    add_text(s, 0.7, 0.85, 12, 0.6, [("Two-column layout", 28, INK, False)])
    add_card(s, 0.7, 1.9, 5.8, 4.7)
    add_card(s, 6.8, 1.9, 5.8, 4.7)
    add_text(s, 1.0, 2.2, 5.2, 0.5, [("Left heading", 18, INK, False)])
    add_text(s, 1.0, 2.8, 5.2, 3.2, [("Body copy. Keep to one idea per column.", 14, INK_2, False)])
    add_text(s, 7.1, 2.2, 5.2, 0.5, [("Right heading", 18, INK, False)])
    add_text(s, 7.1, 2.8, 5.2, 3.2, [("Body copy. Keep to one idea per column.", 14, INK_2, False)])

    # 13 Closing
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_hero)
    add_lockup(s, lockup_dark, 0.7, 2.5, 3.2)
    add_text(s, 0.7, 3.7, 11, 0.5, [("ittdigitalhub.uk", 22, ON_DARK, False)])
    add_text(s, 0.7, 6.7, 11, 0.3, [("Thank you.", 14, ON_DARK_MUTED, False)])

    out = PPT_OUT / "ITT_Digital_Hub_Conference_Starter.pptx"
    prs.save(out)
    print(f"wrote {out.name}")
    return out


def write_readme() -> None:
    text = """# ITT Digital Hub — conference kit

Материали за бъдеща конференция, отделени от `diva-e MARKETING` и преработени по визията на [ittdigitalhub.uk](https://ittdigitalhub.uk).

## Как да ползвате папката

1. **`02-itt-branded/`** — готовите ITT файлове. Започнете оттук.
2. **`01-selected-sources/`** — копие на избраните оригинали от маркетинговата библиотека. Само за справка; не ги показвайте като ITT.

## Готови ITT файлове

| Файл | Предназначение |
|---|---|
| `02-itt-branded/powerpoint/ITT_Digital_Hub_Master.potx` | PowerPoint шаблон (43 лейаута), пренастроен към сайта |
| `02-itt-branded/powerpoint/ITT_Digital_Hub_Master.pptx` | Същият шаблон, отворен като презентация |
| `02-itt-branded/powerpoint/ITT_Digital_Hub_Conference_Starter.pptx` | Стартова колода с копие от сайта (EN) |
| `02-itt-branded/powerpoint/ITT_Poster_1.potx` … `3` | Постер шаблони, пренастроени |
| `02-itt-branded/digital/linkedin-header-1584x396.png` | LinkedIn корица |
| `02-itt-branded/digital/teams-background-1920x1080.png` | Фон за Teams / Zoom |
| `02-itt-branded/print/one-pager-a4.png` | Едностранна листовка |
| `02-itt-branded/print/badge-85x54mm.png` | Бадж (името се редактира) |
| `02-itt-branded/print/rollup-preview-1080x1920.png` | Преглед за рол-ъп |
| `02-itt-branded/icons/` | Икони, преоцветени в ITT navy / signal |
| `02-itt-branded/fonts/` | IBM Plex Sans — инсталирайте ги преди презентация |

Отворете `ITT_Digital_Hub_Master.potx` в PowerPoint: **File → Save as Template**, после **New Slide** ползва ITT лейаутите.

## Визия (от сайта)

- Хартия `#edf0fb`, мастило/marine `#040e31`, сигнал `#002cff`
- Шрифт **IBM Plex Sans** (както на сайта)
- Официални лога от `public/brand/itt-*.png`
- Тъмен hero с navy + signal glow, светли слайдове с меки бели карти

## Какво е избрано от diva-e MARKETING — и защо

**Включено (формат / шаблон, не съдържание на diva-e):**

- PPT master и icon collection
- Постер шаблони и event формати (бадж, рол-ъп, moderation cards)
- LinkedIn header, Teams фон, Word letter templates
- Пет AI колоди само като *структурна* справка (`structure-examples/`)

**Изключено нарочно:**

- Клиентски референции, `Logos Kunden`, case studies (Netto, SMA, Automotive…)
- GTM curriculum, Salesforce/Adobe/SAP at diva-e, credentials / оборот / org chart
- Снимки на служители, employer branding, офиси, песни, kudos, CoC
- Conclusion лога, proprietary шрифтове (GT Planar), Vivid Circle TIFF-ове
- 10 GB Bynder media — не е нужно за ITT колода

Клиентските истории на diva-e не са истории на ITT Digital Hub. Стартовата колода ползва само публичното копие от ittdigitalhub.uk. Липсващи факти не са дописвани.

`structure-examples/` съдържа оригинални diva-e слайдове. **Не ги представяйте като ITT.** Ползвайте ги само за подредба (mission → approach → next steps), после сменете целия текст.

## Какво още да смените преди събитието

- Име на конференцията върху title слайда
- Email / календар — на сайта още е TODO
- Портрет на Иван Томчев, когато има потвърден asset
- Печатни файлове: баджът и рол-ъпът са прегледи; за печат пратете ги на график с bleed

## Повторно генериране

```bash
python3 conference-kit/scripts/build_itt_conference_kit.py
```
"""
    (KIT / "README.md").write_text(text, encoding="utf-8")


def main() -> None:
    if not LIB.exists():
        raise SystemExit(f"Marketing library not found: {LIB}")
    ensure_dirs()
    copy_selected()
    fonts = download_fonts()
    lockup_light = BRAND / "itt-lockup-compact.png"
    lockup_dark = BRAND / "itt-lockup-compact-on-dark.png"

    master_src = SRC_OUT / "powerpoint/ppt_Master_Logo Endorsement.potx"
    if master_src.exists():
        retheme_office(master_src, PPT_OUT / "ITT_Digital_Hub_Master.potx", lockup_light, lockup_dark)
        retheme_office(master_src, PPT_OUT / "ITT_Digital_Hub_Master.pptx", lockup_light, lockup_dark, as_pptx=True)

    icon_potx = SRC_OUT / "powerpoint/Icon Collection.potx"
    if icon_potx.exists():
        retheme_office(icon_potx, PPT_OUT / "ITT_Icon_Collection.potx", lockup_light, lockup_dark)

    for name in [
        "Poster_1.potx",
        "Poster_2.potx",
        "Poster_3.potx",
        "Poster_1_Druck.potx",
        "Poster_2_Druck.potx",
        "Poster_3_Druck.potx",
    ]:
        src = SRC_OUT / "posters" / name
        if src.exists():
            retheme_office(src, PPT_OUT / f"ITT_{name}", lockup_light, lockup_dark)

    for name, out_name in [
        ("Word_Template.dotx", "ITT_Letter_Template.dotx"),
        ("Word_template_Bulgaria_Sofia.dotx", "ITT_Letter_Template_Sofia.dotx"),
    ]:
        src = SRC_OUT / "letter" / name
        if src.exists():
            retheme_office(src, BRANDED / "letter" / out_name, lockup_light, lockup_dark)

    recolor_icons()
    assets = generate_digital_assets(fonts)
    build_starter_deck(assets["hero"], assets["paper"])
    write_readme()
    print("done", KIT)


if __name__ == "__main__":
    main()
